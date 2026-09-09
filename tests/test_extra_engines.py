from __future__ import annotations

from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import fitz
from openpyxl import load_workbook
from PIL import Image
from docx import Document

from papermint_extra_engines import (
    compare_pdfs,
    crop_pdf,
    html_to_pdf,
    images_to_pdf,
    ocr_pdf,
    pdf_to_excel,
    pdf_to_jpg_zip,
    pdf_to_pdfa,
    redact_pdf,
    repair_pdf,
)


class ExtraEngineTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp(prefix="papermint-tests-"))
        self.pdf_a = self.temp_dir / "a.pdf"
        self.pdf_b = self.temp_dir / "b.pdf"
        self._create_pdf(self.pdf_a, "PaperMint SECRET", "Second page")
        self._create_pdf(self.pdf_b, "PaperMint changed", "Second page")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    @staticmethod
    def _create_pdf(path: Path, first: str, second: str):
        document = fitz.open()
        for text in (first, second):
            page = document.new_page(width=595, height=842)
            page.insert_text((72, 100), text, fontsize=16)
            page.insert_text((72, 150), "Name    Amount", fontsize=11)
            page.insert_text((72, 175), "Mint    7 EUR", fontsize=11)
        document.save(path)
        document.close()

    def test_pdf_to_jpg_zip(self):
        output = self.temp_dir / "pages.zip"
        report = pdf_to_jpg_zip(self.pdf_a, output)
        self.assertEqual(report["pages"], 2)
        with ZipFile(output) as archive:
            self.assertEqual(archive.namelist(), ["page-0001.jpg", "page-0002.jpg"])
            self.assertIsNone(archive.testzip())

    def test_images_and_scan_to_pdf(self):
        images = []
        for index, color in enumerate(((255, 220, 200), (200, 220, 255))):
            path = self.temp_dir / f"image-{index}.png"
            Image.new("RGB", (320, 180), color).save(path)
            images.append(path)
        for scan_mode in (False, True):
            output = self.temp_dir / f"images-{scan_mode}.pdf"
            report = images_to_pdf(images, output, scan_mode=scan_mode)
            self.assertEqual(report["pages"], 2)
            with fitz.open(output) as result:
                self.assertEqual(result.page_count, 2)

    def test_html_to_pdf(self):
        source = self.temp_dir / "page.html"
        source.write_text("<h1>PaperMint</h1><p>Safe HTML conversion.</p>", encoding="utf-8")
        output = self.temp_dir / "page.pdf"
        report = html_to_pdf(source, output)
        self.assertGreaterEqual(report["pages"], 1)
        with fitz.open(output) as result:
            self.assertIn("PaperMint", result[0].get_text())

    def test_repair_crop_and_redact(self):
        repaired = self.temp_dir / "repaired.pdf"
        cropped = self.temp_dir / "cropped.pdf"
        redacted = self.temp_dir / "redacted.pdf"
        self.assertEqual(repair_pdf(self.pdf_a, repaired)["pages"], 2)
        original_width = fitz.open(repaired)[0].rect.width
        crop_pdf(repaired, cropped, 10)
        with fitz.open(cropped) as result:
            self.assertLess(result[0].rect.width, original_width)
        report = redact_pdf(self.pdf_a, redacted, "SECRET")
        self.assertEqual(report["redactions"], 1)
        with fitz.open(redacted) as result:
            self.assertNotIn("SECRET", "".join(page.get_text() for page in result))

    def test_compare(self):
        output = self.temp_dir / "comparison.txt"
        report = compare_pdfs([self.pdf_a, self.pdf_b], output)
        self.assertTrue(report["different"])
        content = output.read_text(encoding="utf-8")
        self.assertIn("-PaperMint SECRET", content)
        self.assertIn("+PaperMint changed", content)

    def test_pdf_to_excel(self):
        output = self.temp_dir / "converted.xlsx"
        report = pdf_to_excel(self.pdf_a, output)
        self.assertEqual(report["pages"], 2)
        workbook = load_workbook(output, read_only=True)
        try:
            self.assertEqual(len(workbook.sheetnames), 2)
        finally:
            workbook.close()

    @unittest.skipUnless(shutil.which("tesseract"), "Tesseract is not installed locally")
    def test_ocr_pdf(self):
        output = self.temp_dir / "ocr.docx"
        report = ocr_pdf(self.pdf_a, output)
        self.assertEqual(report["pages"], 2)
        document = Document(output)
        self.assertTrue(document.paragraphs)

    def test_non_english_ocr_language_is_forwarded(self):
        output = self.temp_dir / "ocr-czech.docx"
        with (
            patch("papermint_extra_engines.shutil.which", return_value="/usr/bin/tesseract"),
            patch("papermint_extra_engines.pytesseract.get_languages", return_value=["ces"]),
            patch(
                "papermint_extra_engines.pytesseract.image_to_string",
                return_value="Příliš žluťoučký kůň",
            ) as recognize,
        ):
            report = ocr_pdf(self.pdf_a, output, language="ces")
        self.assertEqual(report["language"], "ces")
        self.assertEqual(recognize.call_args.kwargs["lang"], "ces")
        document = Document(output)
        self.assertIn("Příliš žluťoučký kůň", "\n".join(p.text for p in document.paragraphs))

    @unittest.skipUnless(shutil.which("gs"), "Ghostscript is not installed locally")
    def test_pdfa(self):
        output = self.temp_dir / "archive.pdf"
        report = pdf_to_pdfa(self.pdf_a, output)
        self.assertEqual(report["standard"], "PDF/A-2b")
        with fitz.open(output) as result:
            self.assertEqual(result.page_count, 2)


if __name__ == "__main__":
    unittest.main()
