from pathlib import Path
from zipfile import ZipFile
import json, shutil, tempfile

import fitz
from PIL import Image, ImageDraw
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader

from papermint_merge_engine import merge_pdfs
from papermint_split_engine import split_pdf
from papermint_compress_engine import compress_pdf
from papermint_word_pdf_engine import word_to_pdf
from papermint_rotate_engine import rotate_pdf
from papermint_organize_engine import organize_pdf
from papermint_protect_engine import protect_pdf
from papermint_unlock_engine import unlock_pdf
from papermint_sign_engine import sign_pdf
from papermint_watermark_engine import watermark_pdf
from papermint_page_numbers_engine import add_page_numbers
from papermint_pdf_ppt_engine import pdf_to_powerpoint
from papermint_extra_engines import (
    pdf_to_excel, pdf_to_jpg_zip, office_to_pdf, images_to_pdf, html_to_pdf,
    pdf_to_pdfa, repair_pdf, ocr_pdf, compare_pdfs, redact_pdf, crop_pdf
)
from papermint_v28_engine import convert_v20

ROOT=Path("smoke-output")
ROOT.mkdir(exist_ok=True)
SRC=ROOT/"source.pdf"

doc=fitz.open()
for i in range(1,4):
    p=doc.new_page()
    p.insert_text((72,72),f"PDFaspect smoke test page {i}",fontsize=20)
    p.insert_text((72,110),"SECRET 12345",fontsize=12)
    p.insert_text((72,145),f"Row {i} | Value {i*100}",fontsize=12)
    p.draw_rect(fitz.Rect(65,170,300,250))
doc.save(SRC); doc.close()

# Word sample
wd=Document(); wd.add_heading("PDFaspect Word smoke",0); wd.add_paragraph("Hello from DOCX conversion.")
table=wd.add_table(rows=2,cols=2); table.cell(0,0).text="A"; table.cell(0,1).text="B"; table.cell(1,0).text="1"; table.cell(1,1).text="2"
DOCX=ROOT/"sample.docx"; wd.save(DOCX)

# Excel sample
wb=Workbook(); ws=wb.active; ws.title="Data"; ws.append(["Name","Amount"]); ws.append(["Alpha",123]); ws.append(["Beta",456]); XLSX=ROOT/"sample.xlsx"; wb.save(XLSX); wb.close()

# PPT sample
prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6]); box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1.5)); box.text_frame.text="PDFaspect PowerPoint smoke"; PPTX=ROOT/"sample.pptx"; prs.save(PPTX)

# Image + HTML samples
img=Image.new("RGB",(1200,700),"white"); d=ImageDraw.Draw(img); d.text((80,100),"PDFaspect OCR smoke 98765",fill="black"); JPG=ROOT/"sample.jpg"; img.save(JPG,quality=92)
HTML=ROOT/"sample.html"; HTML.write_text("<html><body><h1>PDFaspect HTML smoke</h1><p>Hello PDF.</p></body></html>",encoding="utf-8")

results={}
def ok(name,path,extra=None):
    p=Path(path); assert p.exists() and p.stat().st_size>0, name
    results[name]={"file":p.name,"bytes":p.stat().st_size,**(extra or {})}

merge=ROOT/"merged.pdf"; r=merge_pdfs([SRC,SRC],merge); assert len(PdfReader(str(merge)).pages)==6; ok("merge",merge,r)
split=ROOT/"split.zip"; r=split_pdf(SRC,split,"1-2,3"); 
with ZipFile(split) as z: assert len(z.namelist())==2 and z.testzip() is None
ok("split",split,r)
comp=ROOT/"compressed.pdf"; r=compress_pdf(SRC,comp); assert len(PdfReader(str(comp)).pages)==3; ok("compress",comp,r)
wp=ROOT/"word-to-pdf.pdf"; r=word_to_pdf(DOCX,wp); assert fitz.open(wp).page_count>=1; ok("word-pdf",wp,r)
rot=ROOT/"rotated.pdf"; r=rotate_pdf(SRC,rot,90); assert PdfReader(str(rot)).pages[0].rotation in (90,270); ok("rotate",rot,r)
org=ROOT/"organized.pdf"; r=organize_pdf(SRC,org,"3,1,2"); assert len(PdfReader(str(org)).pages)==3; ok("organize",org,r)
prot=ROOT/"protected.pdf"; r=protect_pdf(SRC,prot,"smoke123"); pr=PdfReader(str(prot)); assert pr.is_encrypted and pr.decrypt("smoke123"); ok("protect",prot,r)
unl=ROOT/"unlocked.pdf"; r=unlock_pdf(prot,unl,"smoke123"); assert not PdfReader(str(unl)).is_encrypted; ok("unlock",unl,r)
sig=ROOT/"signed.pdf"; r=sign_pdf(SRC,sig,"PDFaspect",1,30,30); assert fitz.open(sig).page_count==3; ok("sign",sig,r)
wm=ROOT/"watermarked.pdf"; r=watermark_pdf(SRC,wm,"PDFASPECT"); assert fitz.open(wm).page_count==3; ok("watermark",wm,r)
num=ROOT/"numbered.pdf"; r=add_page_numbers(SRC,num,1,"bottom-center","page-total",False); assert fitz.open(num).page_count==3; ok("page-numbers",num,r)
ppt=ROOT/"pdf-to-ppt.pptx"; r=pdf_to_powerpoint(SRC,ppt); assert len(Presentation(ppt).slides)==3; ok("pdf-ppt",ppt,r)
xls=ROOT/"pdf-to-excel.xlsx"; r=pdf_to_excel(SRC,xls); chk=load_workbook(xls,read_only=True); assert chk.sheetnames; chk.close(); ok("pdf-excel",xls,r)
jpgzip=ROOT/"pdf-to-jpg.zip"; r=pdf_to_jpg_zip(SRC,jpgzip); 
with ZipFile(jpgzip) as z: assert len(z.namelist())==3 and z.testzip() is None
ok("pdf-jpg",jpgzip,r)
ppdf=ROOT/"ppt-to-pdf.pdf"; r=office_to_pdf(PPTX,ppdf,"ppt-pdf"); assert fitz.open(ppdf).page_count>=1; ok("ppt-pdf",ppdf,r)
epdf=ROOT/"excel-to-pdf.pdf"; r=office_to_pdf(XLSX,epdf,"excel-pdf"); assert fitz.open(epdf).page_count>=1; ok("excel-pdf",epdf,r)
ipdf=ROOT/"jpg-to-pdf.pdf"; r=images_to_pdf([JPG],ipdf,False); assert fitz.open(ipdf).page_count==1; ok("jpg-pdf",ipdf,r)
scan=ROOT/"scan-to-pdf.pdf"; r=images_to_pdf([JPG],scan,True); assert fitz.open(scan).page_count==1; ok("scan-pdf",scan,r)
hpdf=ROOT/"html-to-pdf.pdf"; r=html_to_pdf(HTML,hpdf); assert fitz.open(hpdf).page_count>=1; ok("html-pdf",hpdf,r)
pdfa=ROOT/"pdfa.pdf"; r=pdf_to_pdfa(SRC,pdfa); assert fitz.open(pdfa).page_count==3; ok("pdfa",pdfa,r)
rep=ROOT/"repaired.pdf"; r=repair_pdf(SRC,rep); assert fitz.open(rep).page_count==3; ok("repair",rep,r)
ocr=ROOT/"ocr.docx"; r=ocr_pdf(scan,ocr,"eng"); od=Document(ocr); assert len(od.paragraphs)>=1; ok("ocr",ocr,r)
cmp=ROOT/"comparison.txt"; r=compare_pdfs([SRC,rot],cmp); assert cmp.read_text(encoding="utf-8").strip(); ok("compare",cmp,r)
red=ROOT/"redacted.pdf"; r=redact_pdf(SRC,red,"SECRET 12345"); rd=fitz.open(red); assert "SECRET 12345" not in "".join(p.get_text() for p in rd); rd.close(); ok("redact",red,r)
crop=ROOT/"cropped.pdf"; r=crop_pdf(SRC,crop,5.0); assert fitz.open(crop).page_count==3; ok("crop",crop,r)

# Real PDF -> Word engine conversion (production-style QA disabled)
p2w=ROOT/"pdf-to-word.docx"; r=convert_v20(SRC,p2w,ROOT/"pdf-word-work",qa=False); Document(p2w); ok("pdf-word",p2w,r)

(ROOT/"results.json").write_text(json.dumps(results,indent=2,ensure_ascii=False),encoding="utf-8")
print(json.dumps({"passed":len(results),"tools":sorted(results)},ensure_ascii=False))
