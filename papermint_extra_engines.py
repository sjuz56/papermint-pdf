"""Conversion and document-processing engines for PaperMint's remaining tools."""

from __future__ import annotations

import difflib
from io import BytesIO
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
from urllib.parse import urlparse
from zipfile import ZIP_DEFLATED, ZipFile

import fitz
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageEnhance, ImageOps, UnidentifiedImageError
import pdfplumber
import pytesseract
from weasyprint import HTML


MAX_PDF_PAGES = 500
MAX_IMAGES = 100
MAX_IMAGE_PIXELS = 50_000_000
MAX_RASTER_OUTPUT_BYTES = 250 * 1024 * 1024
Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS


class ExtraToolError(RuntimeError):
    """A user-safe failure raised by one of the extra tools."""


class PdfExcelError(ExtraToolError):
    pass


class PdfJpgError(ExtraToolError):
    pass


class OfficePdfError(ExtraToolError):
    pass


class ImagePdfError(ExtraToolError):
    pass


class HtmlPdfError(ExtraToolError):
    pass


class PdfMaintenanceError(ExtraToolError):
    pass


class OcrError(ExtraToolError):
    pass


class CompareError(ExtraToolError):
    pass


def _atomic_temp(destination: Path, suffix: str) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(
        prefix=f".{destination.stem}-",
        suffix=suffix,
        dir=destination.parent,
    )
    os.close(fd)
    return Path(name)


def _publish(temporary: Path, destination: Path) -> None:
    os.replace(temporary, destination)


def _open_pdf(source: Path, error_type: type[ExtraToolError]) -> fitz.Document:
    try:
        document = fitz.open(source)
    except Exception as exc:
        raise error_type("The uploaded file is not a readable PDF.") from exc
    if not document.is_pdf or document.page_count < 1:
        document.close()
        raise error_type("The PDF does not contain any readable pages.")
    if document.needs_pass:
        document.close()
        raise error_type("Unlock the PDF before using this tool.")
    if document.page_count > MAX_PDF_PAGES:
        document.close()
        raise error_type(f"The PDF can contain at most {MAX_PDF_PAGES} pages.")
    return document


def _validate_pdf_output(path: Path, expected_pages: int | None = None) -> int:
    try:
        with fitz.open(path) as document:
            if not document.is_pdf or document.needs_pass or document.page_count < 1:
                raise ValueError
            if expected_pages is not None and document.page_count != expected_pages:
                raise ValueError
            document[0].get_pixmap(matrix=fitz.Matrix(0.2, 0.2), alpha=False)
            return document.page_count
    except Exception as exc:
        raise PdfMaintenanceError("The created PDF failed its integrity check.") from exc


def pdf_to_excel(source: str | Path, output: str | Path) -> dict:
    source_path = Path(source)
    destination = Path(output)
    check = _open_pdf(source_path, PdfExcelError)
    page_count = check.page_count
    check.close()
    temporary = _atomic_temp(destination, ".xlsx")
    sheets = 0
    try:
        workbook = Workbook()
        workbook.remove(workbook.active)
        with pdfplumber.open(source_path) as pdf:
            for page_number, page in enumerate(pdf.pages, 1):
                tables = page.extract_tables() or []
                if tables:
                    for table_number, table in enumerate(tables, 1):
                        sheet = workbook.create_sheet(f"P{page_number} T{table_number}"[:31])
                        sheets += 1
                        for row_number, row in enumerate(table or [], 1):
                            for column_number, value in enumerate(row or [], 1):
                                sheet.cell(row_number, column_number, value)
                else:
                    sheet = workbook.create_sheet(f"Page {page_number}"[:31])
                    sheets += 1
                    for row_number, line in enumerate((page.extract_text() or "").splitlines(), 1):
                        sheet.cell(row_number, 1, line)
        if sheets == 0:
            workbook.create_sheet("Document")
            sheets = 1
        workbook.save(temporary)
        from openpyxl import load_workbook

        validation = load_workbook(temporary, read_only=True)
        try:
            if not validation.sheetnames:
                raise PdfExcelError("The Excel workbook failed its integrity check.")
        finally:
            validation.close()
        _publish(temporary, destination)
        return {"pages": page_count, "sheets": sheets}
    except PdfExcelError:
        raise
    except Exception as exc:
        raise PdfExcelError("The PDF could not be converted to Excel.") from exc
    finally:
        temporary.unlink(missing_ok=True)


def pdf_to_jpg_zip(source: str | Path, output: str | Path, dpi: int = 160) -> dict:
    destination = Path(output)
    document = _open_pdf(Path(source), PdfJpgError)
    temporary = _atomic_temp(destination, ".zip")
    try:
        scale = max(1.0, min(4.0, dpi / 72.0))
        with ZipFile(temporary, "w", compression=ZIP_DEFLATED, compresslevel=6) as archive:
            rendered_bytes = 0
            for page_number, page in enumerate(document, 1):
                longest = max(page.rect.width, page.rect.height)
                page_scale = min(scale, 3_500 / longest) if longest > 0 else scale
                pixmap = page.get_pixmap(matrix=fitz.Matrix(page_scale, page_scale), alpha=False)
                jpg = pixmap.tobytes("jpeg", jpg_quality=88)
                rendered_bytes += len(jpg)
                if rendered_bytes > MAX_RASTER_OUTPUT_BYTES:
                    raise PdfJpgError("The rendered JPG files exceed the 250 MB output limit.")
                archive.writestr(
                    f"page-{page_number:04d}.jpg",
                    jpg,
                )
        with ZipFile(temporary) as archive:
            if len(archive.namelist()) != document.page_count or archive.testzip() is not None:
                raise PdfJpgError("The JPG archive failed its integrity check.")
        pages = document.page_count
        _publish(temporary, destination)
        return {"pages": pages, "dpi": dpi}
    except PdfJpgError:
        raise
    except Exception as exc:
        raise PdfJpgError("The PDF pages could not be converted to JPG.") from exc
    finally:
        document.close()
        temporary.unlink(missing_ok=True)


OFFICE_SUFFIXES = {
    "ppt-pdf": {".ppt", ".pptx"},
    "excel-pdf": {".xls", ".xlsx"},
}


def office_to_pdf(source: str | Path, output: str | Path, tool: str) -> dict:
    source_path = Path(source)
    destination = Path(output)
    allowed = OFFICE_SUFFIXES.get(tool)
    if not allowed or source_path.suffix.lower() not in allowed:
        label = "PowerPoint" if tool == "ppt-pdf" else "Excel"
        raise OfficePdfError(f"{label} to PDF received an unsupported file type.")
    soffice = os.environ.get("PAPERMINT_SOFFICE", "").strip()
    soffice = soffice if soffice and Path(soffice).is_file() else shutil.which("libreoffice") or shutil.which("soffice")
    if not soffice:
        raise OfficePdfError("The document converter is not available on this server.")

    destination.parent.mkdir(parents=True, exist_ok=True)
    work_root = Path(tempfile.mkdtemp(prefix="papermint-office-pdf-", dir=destination.parent))
    profile = work_root / "profile"
    converted_dir = work_root / "output"
    home_dir = work_root / "home"
    for directory in (profile, converted_dir, home_dir):
        directory.mkdir(parents=True, exist_ok=True)
    local_source = work_root / ("document" + source_path.suffix.lower())
    temporary = _atomic_temp(destination, ".pdf")
    try:
        shutil.copyfile(source_path, local_source)
        environment = os.environ.copy()
        environment.update({"HOME": str(home_dir), "XDG_CACHE_HOME": str(home_dir / ".cache")})
        command = [
            str(soffice), "--headless", "--nologo", "--nodefault", "--nofirststartwizard",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--convert-to", "pdf", "--outdir", str(converted_dir), str(local_source),
        ]
        try:
            completed = subprocess.run(
                command, env=environment, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                text=True, timeout=180, check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise OfficePdfError("The document conversion timed out.") from exc
        candidate = converted_dir / "document.pdf"
        if completed.returncode != 0 or not candidate.is_file():
            raise OfficePdfError("The server could not convert this document to PDF.")
        shutil.copyfile(candidate, temporary)
        pages = _validate_pdf_output(temporary)
        _publish(temporary, destination)
        return {"pages": pages, "converter": "libreoffice"}
    except OfficePdfError:
        raise
    except PdfMaintenanceError as exc:
        raise OfficePdfError(str(exc)) from exc
    except Exception as exc:
        raise OfficePdfError("The document could not be converted to PDF.") from exc
    finally:
        temporary.unlink(missing_ok=True)
        shutil.rmtree(work_root, ignore_errors=True)


def _load_image(path: Path, scan_mode: bool) -> Image.Image:
    try:
        with Image.open(path) as opened:
            opened.load()
            image = ImageOps.exif_transpose(opened).convert("RGB")
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
        raise ImagePdfError("One of the uploaded images is invalid or too large.") from exc
    if scan_mode:
        gray = ImageOps.grayscale(image)
        gray = ImageOps.autocontrast(gray, cutoff=1)
        gray = ImageEnhance.Contrast(gray).enhance(1.12)
        image = gray.convert("RGB")
    return image


def images_to_pdf(sources: list[str] | list[Path], output: str | Path, scan_mode: bool = False) -> dict:
    if not sources:
        raise ImagePdfError("Upload at least one image.")
    if len(sources) > MAX_IMAGES:
        raise ImagePdfError(f"You can combine up to {MAX_IMAGES} images at once.")
    destination = Path(output)
    temporary = _atomic_temp(destination, ".pdf")
    images: list[Image.Image] = []
    try:
        images = [_load_image(Path(path), scan_mode) for path in sources]
        images[0].save(
            temporary,
            "PDF",
            save_all=True,
            append_images=images[1:],
            resolution=150.0,
        )
        pages = _validate_pdf_output(temporary, len(images))
        _publish(temporary, destination)
        return {"pages": pages, "scan_enhancement": scan_mode}
    except ImagePdfError:
        raise
    except PdfMaintenanceError as exc:
        raise ImagePdfError(str(exc)) from exc
    except Exception as exc:
        raise ImagePdfError("The images could not be converted to PDF.") from exc
    finally:
        for image in images:
            image.close()
        temporary.unlink(missing_ok=True)


def _safe_url_fetcher(url: str, *args, **kwargs):
    scheme = urlparse(url).scheme.lower()
    if scheme == "data":
        from weasyprint.urls import default_url_fetcher

        return default_url_fetcher(url, *args, **kwargs)
    raise ValueError("External resources are disabled for HTML conversion.")


def html_to_pdf(source: str | Path, output: str | Path) -> dict:
    source_path = Path(source)
    destination = Path(output)
    if source_path.suffix.lower() not in {".html", ".htm"}:
        raise HtmlPdfError("HTML to PDF accepts HTML and HTM files only.")
    if not source_path.is_file() or source_path.stat().st_size == 0:
        raise HtmlPdfError("The uploaded HTML file is empty.")
    if source_path.stat().st_size > 5 * 1024 * 1024:
        raise HtmlPdfError("The HTML file can be up to 5 MB.")
    temporary = _atomic_temp(destination, ".pdf")
    try:
        markup = source_path.read_text(encoding="utf-8", errors="replace")
        HTML(string=markup, url_fetcher=_safe_url_fetcher).write_pdf(temporary)
        pages = _validate_pdf_output(temporary)
        _publish(temporary, destination)
        return {"pages": pages, "external_resources": "blocked"}
    except PdfMaintenanceError as exc:
        raise HtmlPdfError(str(exc)) from exc
    except Exception as exc:
        raise HtmlPdfError("The HTML file could not be converted to PDF.") from exc
    finally:
        temporary.unlink(missing_ok=True)


def repair_pdf(source: str | Path, output: str | Path) -> dict:
    destination = Path(output)
    document = _open_pdf(Path(source), PdfMaintenanceError)
    temporary = _atomic_temp(destination, ".pdf")
    try:
        pages = document.page_count
        document.save(temporary, garbage=4, deflate=True, clean=True)
        document.close()
        document = None
        _validate_pdf_output(temporary, pages)
        _publish(temporary, destination)
        return {"pages": pages, "rewritten": True}
    except PdfMaintenanceError:
        raise
    except Exception as exc:
        raise PdfMaintenanceError("The PDF could not be repaired.") from exc
    finally:
        if document is not None:
            document.close()
        temporary.unlink(missing_ok=True)


def pdf_to_pdfa(source: str | Path, output: str | Path) -> dict:
    source_path = Path(source)
    destination = Path(output)
    check = _open_pdf(source_path, PdfMaintenanceError)
    pages = check.page_count
    check.close()
    ghostscript = shutil.which("gs")
    if not ghostscript:
        raise PdfMaintenanceError("PDF/A conversion is not available on this server.")
    icc_candidates = [
        Path("/usr/share/color/icc/ghostscript/srgb.icc"),
        Path("/usr/share/color/icc/sRGB.icc"),
        Path("/usr/share/ghostscript/iccprofiles/default_rgb.icc"),
    ]
    icc = next((path for path in icc_candidates if path.is_file()), None)
    if icc is None:
        raise PdfMaintenanceError("The PDF/A color profile is missing on this server.")
    temporary = _atomic_temp(destination, ".pdf")
    definition = _atomic_temp(destination, ".ps")
    try:
        icc_ps = str(icc).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        definition.write_text(
            "[/_objdef {icc_PDFA} /type /stream /OBJ pdfmark\n"
            "[{icc_PDFA} << /N 3 >> /PUT pdfmark\n"
            f"[{{icc_PDFA}} ({icc_ps}) (r) file /PUT pdfmark\n"
            "[/_objdef {OutputIntent_PDFA} /type /dict /OBJ pdfmark\n"
            "[{OutputIntent_PDFA} << /Type /OutputIntent /S /GTS_PDFA1 "
            "/DestOutputProfile {icc_PDFA} /OutputConditionIdentifier (sRGB) >> /PUT pdfmark\n"
            "[{Catalog} << /OutputIntents [ {OutputIntent_PDFA} ] >> /PUT pdfmark\n",
            encoding="ascii",
        )
        command = [
            ghostscript, "-dPDFA=2", "-dBATCH", "-dNOPAUSE", "-dSAFER",
            f"--permit-file-read={icc}",
            "-sDEVICE=pdfwrite", "-dPDFACompatibilityPolicy=1",
            "-sColorConversionStrategy=RGB", "-dProcessColorModel=/DeviceRGB",
            f"-sOutputFile={temporary}", str(definition), str(source_path),
        ]
        try:
            completed = subprocess.run(
                command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                text=True, timeout=240, check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise PdfMaintenanceError("PDF/A conversion timed out.") from exc
        if completed.returncode != 0 or not temporary.is_file():
            raise PdfMaintenanceError("The PDF could not be converted to PDF/A.")
        _validate_pdf_output(temporary, pages)
        _publish(temporary, destination)
        return {"pages": pages, "standard": "PDF/A-2b"}
    finally:
        temporary.unlink(missing_ok=True)
        definition.unlink(missing_ok=True)


def redact_pdf(source: str | Path, output: str | Path, text: str) -> dict:
    needle = text.strip()
    if not needle:
        raise PdfMaintenanceError("Enter text to redact.")
    if len(needle) > 200:
        raise PdfMaintenanceError("Redaction text can contain at most 200 characters.")
    destination = Path(output)
    document = _open_pdf(Path(source), PdfMaintenanceError)
    temporary = _atomic_temp(destination, ".pdf")
    matches = 0
    try:
        pages = document.page_count
        for page in document:
            rectangles = page.search_for(needle)
            matches += len(rectangles)
            for rectangle in rectangles:
                page.add_redact_annot(rectangle, fill=(0, 0, 0))
            if rectangles:
                page.apply_redactions()
        if matches == 0:
            raise PdfMaintenanceError("The specified text was not found in the PDF.")
        document.save(temporary, garbage=4, deflate=True, clean=True)
        document.close()
        document = None
        _validate_pdf_output(temporary, pages)
        _publish(temporary, destination)
        return {"pages": pages, "redactions": matches}
    except PdfMaintenanceError:
        raise
    except Exception as exc:
        raise PdfMaintenanceError("The text could not be redacted from this PDF.") from exc
    finally:
        if document is not None:
            document.close()
        temporary.unlink(missing_ok=True)


def crop_pdf(source: str | Path, output: str | Path, margin_mm: float) -> dict:
    if margin_mm < 0 or margin_mm > 100:
        raise PdfMaintenanceError("Crop margin must be between 0 and 100 mm.")
    destination = Path(output)
    document = _open_pdf(Path(source), PdfMaintenanceError)
    temporary = _atomic_temp(destination, ".pdf")
    margin = margin_mm * 72.0 / 25.4
    try:
        pages = document.page_count
        for page in document:
            box = page.cropbox
            if box.width <= 2 * margin + 1 or box.height <= 2 * margin + 1:
                raise PdfMaintenanceError("The crop margin is too large for at least one page.")
            page.set_cropbox(
                fitz.Rect(box.x0 + margin, box.y0 + margin, box.x1 - margin, box.y1 - margin)
            )
        document.save(temporary, garbage=4, deflate=True, clean=True)
        document.close()
        document = None
        _validate_pdf_output(temporary, pages)
        _publish(temporary, destination)
        return {"pages": pages, "margin_mm": margin_mm}
    except PdfMaintenanceError:
        raise
    except Exception as exc:
        raise PdfMaintenanceError("The PDF pages could not be cropped.") from exc
    finally:
        if document is not None:
            document.close()
        temporary.unlink(missing_ok=True)


def ocr_pdf(source: str | Path, output: str | Path, language: str = "eng") -> dict:
    if not shutil.which("tesseract"):
        raise OcrError("OCR is not available on this server.")
    destination = Path(output)
    document = _open_pdf(Path(source), OcrError)
    if document.page_count > 100:
        document.close()
        raise OcrError("OCR accepts PDFs with up to 100 pages.")
    temporary = _atomic_temp(destination, ".docx")
    try:
        result = Document()
        for page_number, page in enumerate(document):
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2.2, 2.2), alpha=False)
            image = Image.open(BytesIO(pixmap.tobytes("png")))
            recognized = pytesseract.image_to_string(image, lang=language).strip()
            result.add_paragraph(recognized)
            if page_number < document.page_count - 1:
                result.add_page_break()
        result.save(temporary)
        validation = Document(temporary)
        if not validation.paragraphs:
            raise OcrError("The OCR document failed its integrity check.")
        pages = document.page_count
        _publish(temporary, destination)
        return {"pages": pages, "language": language}
    except OcrError:
        raise
    except Exception as exc:
        raise OcrError("Text recognition failed for this PDF.") from exc
    finally:
        document.close()
        temporary.unlink(missing_ok=True)


def compare_pdfs(sources: list[str] | list[Path], output: str | Path) -> dict:
    if len(sources) != 2:
        raise CompareError("Upload exactly two PDFs to compare.")
    documents: list[fitz.Document] = []
    try:
        for path in sources:
            documents.append(_open_pdf(Path(path), CompareError))
    except Exception:
        for document in documents:
            document.close()
        raise
    destination = Path(output)
    temporary = _atomic_temp(destination, ".txt")
    try:
        left = "\n".join(page.get_text("text") for page in documents[0])
        right = "\n".join(page.get_text("text") for page in documents[1])
        diff_lines = list(
            difflib.unified_diff(
                left.splitlines(), right.splitlines(),
                fromfile="PDF A", tofile="PDF B", lineterm="",
            )
        )
        content = "\n".join(diff_lines)
        if not content:
            content = "No text differences found.\n"
        temporary.write_text(content, encoding="utf-8")
        _publish(temporary, destination)
        return {
            "left_pages": documents[0].page_count,
            "right_pages": documents[1].page_count,
            "different": bool(diff_lines),
        }
    except CompareError:
        raise
    except Exception as exc:
        raise CompareError("The PDFs could not be compared.") from exc
    finally:
        for document in documents:
            document.close()
        temporary.unlink(missing_ok=True)
