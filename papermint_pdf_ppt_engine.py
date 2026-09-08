"""Layout-faithful PDF to PowerPoint conversion for PaperMint."""

from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
import tempfile

import fitz
from pptx import Presentation
from pptx.util import Inches


MAX_PAGES = 300
RENDER_DPI = 160
MAX_RENDER_DIMENSION = 3_500
SLIDE_LONG_SIDE_INCHES = 13.333
MIN_SLIDE_SHORT_SIDE_INCHES = 1.0


class PdfPowerPointError(RuntimeError):
    """A PDF could not be converted to PowerPoint."""


def _slide_size(page_width: float, page_height: float) -> tuple[int, int]:
    if page_width <= 0 or page_height <= 0:
        raise PdfPowerPointError("The first PDF page has an invalid size.")

    if page_width >= page_height:
        width_inches = SLIDE_LONG_SIDE_INCHES
        height_inches = max(
            MIN_SLIDE_SHORT_SIDE_INCHES,
            width_inches * page_height / page_width,
        )
    else:
        height_inches = SLIDE_LONG_SIDE_INCHES
        width_inches = max(
            MIN_SLIDE_SHORT_SIDE_INCHES,
            height_inches * page_width / page_height,
        )

    return Inches(width_inches), Inches(height_inches)


def _render_page(page: fitz.Page) -> tuple[BytesIO, int, int]:
    scale = RENDER_DPI / 72.0
    longest_dimension = max(page.rect.width, page.rect.height)
    if longest_dimension * scale > MAX_RENDER_DIMENSION:
        scale = MAX_RENDER_DIMENSION / longest_dimension

    pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
    return BytesIO(pixmap.tobytes("png")), pixmap.width, pixmap.height


def _fit_image(
    image_width: int,
    image_height: int,
    slide_width: int,
    slide_height: int,
) -> tuple[int, int, int, int]:
    if image_width <= 0 or image_height <= 0:
        raise PdfPowerPointError("A PDF page could not be rendered.")

    scale = min(slide_width / image_width, slide_height / image_height)
    width = max(1, round(image_width * scale))
    height = max(1, round(image_height * scale))
    left = (slide_width - width) // 2
    top = (slide_height - height) // 2
    return left, top, width, height


def pdf_to_powerpoint(source: str | Path, output: str | Path) -> dict:
    """Render every PDF page as one sharp, undistorted PowerPoint slide."""
    source_path = Path(source)
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document: fitz.Document | None = None
    temporary_path: Path | None = None

    try:
        try:
            document = fitz.open(source_path)
        except Exception as exc:
            raise PdfPowerPointError("The uploaded file is not a readable PDF.") from exc

        if document.needs_pass:
            raise PdfPowerPointError("Unlock the PDF before converting it to PowerPoint.")

        page_count = document.page_count
        if page_count < 1:
            raise PdfPowerPointError("The PDF does not contain any pages.")
        if page_count > MAX_PAGES:
            raise PdfPowerPointError(f"The PDF can contain at most {MAX_PAGES} pages.")

        first_page = document[0]
        presentation = Presentation()
        presentation.slide_width, presentation.slide_height = _slide_size(
            first_page.rect.width,
            first_page.rect.height,
        )
        blank_layout = presentation.slide_layouts[6]

        for page in document:
            rendered, pixmap_width, pixmap_height = _render_page(page)
            left, top, width, height = _fit_image(
                pixmap_width,
                pixmap_height,
                presentation.slide_width,
                presentation.slide_height,
            )
            slide = presentation.slides.add_slide(blank_layout)
            slide.shapes.add_picture(rendered, left, top, width=width, height=height)

        fd, temp_name = tempfile.mkstemp(
            prefix=f".{output_path.stem}-",
            suffix=".pptx",
            dir=output_path.parent,
        )
        os.close(fd)
        temporary_path = Path(temp_name)
        presentation.save(temporary_path)

        check = Presentation(temporary_path)
        if len(check.slides) != page_count:
            raise PdfPowerPointError(
                "The PowerPoint failed its final integrity check."
            )

        os.replace(temporary_path, output_path)
        temporary_path = None

        return {
            "pages": page_count,
            "slides": page_count,
            "mode": "page_images",
            "render_dpi": RENDER_DPI,
        }
    except PdfPowerPointError:
        raise
    except Exception as exc:
        raise PdfPowerPointError(
            "The PDF could not be converted to PowerPoint."
        ) from exc
    finally:
        if document is not None:
            document.close()
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
